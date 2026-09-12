#!/usr/bin/env python3
"""Собирает из книги простой файл для чтения: артикул, название, теги, описание.

Рабочая книга широкая — 36 колонок на листе «Товары» плюс одиннадцать листов
категорий. Тексты в ней есть, но найти их глазами трудно: название в колонке C,
хештеги в X, описание в Y. Этот файл — та же информация на одном экране.

В Ozon он не заливается. Только чтобы прочитать и поправить.

Запуск:
    python3 scripts/ozon_export_texts.py docs/ozon/Озон_рабочая_книга.xlsx docs/ozon/Озон_тексты.xlsx
"""

import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("Артикул", "Артикул *", 18),
    ("Категория", "Категория *", 24),
    ("Название товара", "Название товара", 55),
    ("Хештеги", "#Хештеги", 45),
    ("Описание", "Аннотация", 95),
]


def main(src, dst):
    book = openpyxl.load_workbook(src)
    ws = book["Товары"]
    header = [c.value for c in ws[1]]
    idx = {title: header.index(title) + 1 for _, title, _ in COLUMNS}

    out = openpyxl.Workbook()
    sheet = out.active
    sheet.title = "Тексты"

    head_fill = PatternFill("solid", fgColor="D9E1F2")
    for i, (label, _, width) in enumerate(COLUMNS, 1):
        cell = sheet.cell(row=1, column=i, value=label)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "C2"

    missing = PatternFill("solid", fgColor="FCE4D6")
    row_out = 2
    for r in range(3, ws.max_row + 1):
        article = ws.cell(row=r, column=idx["Артикул *"]).value
        if not article:
            continue
        for i, (_, title, _) in enumerate(COLUMNS, 1):
            value = ws.cell(row=r, column=idx[title]).value
            cell = sheet.cell(row=row_out, column=i, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if value in (None, "") and title in ("Название товара", "Аннотация"):
                cell.fill = missing
        sheet.row_dimensions[row_out].height = 90
        row_out += 1

    out.save(dst)
    print(f"строк: {row_out - 2}  →  {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
