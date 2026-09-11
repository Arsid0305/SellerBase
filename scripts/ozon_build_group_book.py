#!/usr/bin/env python3
"""Собирает из рабочей книги Ozon книгу по группам — как книга WB.

Рабочая книга устроена под загрузку: один широкий лист «Товары» на все
категории плюс одиннадцать листов категорий с особыми характеристиками.
Читать её глазами неудобно — тексты стоят в колонках C, X и Y.

Здесь то же самое, но по-человечески: лист на группу, в строке товар,
в колонках всё, что к нему относится. Заполненное подсвечено зелёным,
незаполненное обязательное — оранжевым, решения владелицы — серым.

В Ozon эта книга не заливается. Для работы и чтения.

Запуск:
    python3 scripts/ozon_build_group_book.py \\
        docs/ozon/Озон_рабочая_книга.xlsx docs/ozon/Озон_по_группам.xlsx
"""

import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREEN = PatternFill("solid", fgColor="E2EFDA")   # заполнено
ORANGE = PatternFill("solid", fgColor="FCE4D6")  # обязательное, пусто
GREY = PatternFill("solid", fgColor="EDEDED")    # ждёт решения владелицы
HEAD = PatternFill("solid", fgColor="D9E1F2")
HEAD_REQ = PatternFill("solid", fgColor="F8CBAD")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Колонки листа «Товары», которые несут смысл для работы. Порядок — рабочий:
# сначала тексты, потом склад и логистика, потом то, что ждёт владелицу.
FROM_GOODS = [
    ("Название товара", 46, "текст"),
    ("#Хештеги", 38, "текст"),
    ("Аннотация", 90, "текст"),
    ("Штрихкод (Серийный номер / EAN)", 18, "факт"),
    ("Бренд *", 12, "факт"),
    ("Вес в упаковке, г *", 11, "факт"),
    ("Длина упаковки, мм *", 11, "факт"),
    ("Ширина упаковки, мм *", 11, "факт"),
    ("Высота упаковки, мм *", 11, "факт"),
    ("Предельная цена без акций, руб. *", 13, "владелица"),
    ("НДС, % *", 9, "владелица"),
    ("Ссылка на главное фото *", 22, "владелица"),
    ("Цвет товара", 16, "владелица"),
    ("Страна-изготовитель", 16, "владелица"),
]

RENAME = {
    "Аннотация": "Описание",
    "Штрихкод (Серийный номер / EAN)": "Штрихкод",
    "Бренд *": "Бренд",
    "Вес в упаковке, г *": "Вес, г",
    "Длина упаковки, мм *": "Длина, мм",
    "Ширина упаковки, мм *": "Ширина, мм",
    "Высота упаковки, мм *": "Высота, мм",
    "Предельная цена без акций, руб. *": "Цена, руб",
    "НДС, % *": "НДС, %",
    "Ссылка на главное фото *": "Фото (ссылка)",
}


def read_goods(book):
    ws = book["Товары"]
    header = [c.value for c in ws[1]]
    rows = []
    for r in range(3, ws.max_row + 1):
        values = {header[i]: ws.cell(row=r, column=i + 1).value for i in range(len(header))}
        if values.get("Артикул *"):
            rows.append(values)
    return rows


def read_category(book, name):
    """Лист категории: строка 1 — заголовки, 2 — подсказки, данные с 3."""
    if name not in book.sheetnames:
        return [], {}
    ws = book[name]
    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    by_article = {}
    for r in range(3, ws.max_row + 1):
        article = ws.cell(row=r, column=1).value
        if not article:
            continue
        by_article[article] = {
            header[c - 1]: ws.cell(row=r, column=c).value
            for c in range(2, ws.max_column + 1)
        }
    return [h for h in header[1:] if h], by_article


def write_group(out, group, goods, cat_columns, cat_values):
    sheet = out.create_sheet(group[:31])

    columns = [("Артикул", 15, "факт")]
    columns += [(RENAME.get(t, t), w, kind) for t, w, kind in FROM_GOODS]
    columns += [(c, 20, "категория") for c in cat_columns]

    for i, (label, width, kind) in enumerate(columns, 1):
        cell = sheet.cell(row=1, column=i, value=label)
        cell.font = Font(bold=True, size=10)
        cell.fill = HEAD_REQ if kind in ("владелица", "категория") else HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "B2"

    for row_i, item in enumerate(goods, 2):
        article = item["Артикул *"]
        values = [article]
        values += [item.get(t) for t, _, _ in FROM_GOODS]
        values += [cat_values.get(article, {}).get(c) for c in cat_columns]

        for i, ((_, _, kind), value) in enumerate(zip(columns, values), 1):
            cell = sheet.cell(row=row_i, column=i, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            if value not in (None, ""):
                cell.fill = GREEN
            elif kind == "владелица":
                cell.fill = GREY
            elif kind in ("текст", "категория"):
                cell.fill = ORANGE
        sheet.row_dimensions[row_i].height = 110
    return len(goods)


def write_intro(out, stats):
    sheet = out.create_sheet("Как читать", 0)
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["E"].width = 58

    lines = [
        ("Книга Ozon по группам", None, None, None, None),
        ("", None, None, None, None),
        ("Лист на группу. В строке товар, в колонках всё про него.", None, None, None, None),
        ("Зелёное — заполнено. Оранжевое — надо заполнить. Серое — ждёт вашего решения.",
         None, None, None, None),
        ("", None, None, None, None),
        ("Группа", "Товаров", "Названий", "Описаний", "Что осталось"),
    ]
    for i, row in enumerate(lines, 1):
        for j, value in enumerate(row, 1):
            cell = sheet.cell(row=i, column=j, value=value)
            if i == 1:
                cell.font = Font(bold=True, size=14)
            if i == 6:
                cell.font = Font(bold=True)
                cell.fill = HEAD

    row = 7
    for group, total, named, described, rest in stats:
        for j, value in enumerate((group, total, named, described, rest), 1):
            cell = sheet.cell(row=row, column=j, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if j == 4 and described < total:
                cell.fill = ORANGE
            elif j == 4:
                cell.fill = GREEN
        sheet.row_dimensions[row].height = 30
        row += 1

    row += 1
    for text in (
        "Что ещё не сделано по всей книге:",
        "1. Характеристики категорий — оранжевые колонки справа на каждом листе."
        " Они дают 55 баллов контент-рейтинга из 100, это следующая работа.",
        "2. Цена, НДС, фото, цвет и страна — серые колонки. Ваше решение и ваши данные.",
        "3. Инфографика, видео и rich-контент — не разбираем пока.",
        "",
        "Эта книга в Ozon не заливается. Для работы и чтения.",
    ):
        sheet.cell(row=row, column=1, value=text).alignment = Alignment(wrap_text=True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        sheet.row_dimensions[row].height = 30
        row += 1


def main(src, dst):
    book = openpyxl.load_workbook(src)
    goods = read_goods(book)

    groups = {}
    for item in goods:
        groups.setdefault(item.get("Категория *") or "Без категории", []).append(item)

    out = openpyxl.Workbook()
    out.remove(out.active)

    stats = []
    for group in sorted(groups, key=lambda g: (g == "Без категории", -len(groups[g]))):
        items = groups[group]
        cat_columns, cat_values = read_category(book, group)
        # ТН ВЭД вперёд. Имя колонки берём как есть: у «Кухонного ножа»
        # оно без звёздочки, у остальных со звёздочкой.
        tnved = [c for c in cat_columns if str(c).startswith("ТН ВЭД")]
        cat_columns = tnved + [c for c in cat_columns if c not in tnved]
        write_group(out, group, items, cat_columns, cat_values)

        named = sum(1 for i in items if i.get("Название товара"))
        described = sum(1 for i in items if i.get("Аннотация"))
        empty_cat = [
            c for c in cat_columns
            if not any(cat_values.get(i["Артикул *"], {}).get(c) for i in items)
        ]
        rest = ", ".join(empty_cat) if empty_cat else "характеристики заполнены"
        stats.append((group, len(items), named, described, rest))

    write_intro(out, stats)
    out.save(dst)

    total = sum(s[1] for s in stats)
    print(f"групп {len(stats)}, товаров {total}  →  {dst}")
    for group, n, named, described, _ in stats:
        print(f"  {group:<32} товаров {n:>2}, названий {named:>2}, описаний {described:>2}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
