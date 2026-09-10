#!/usr/bin/env python3
"""Проверяет хештеги Ozon и кладёт их в книгу по категориям.

Правила Ozon (docs/ozon/ozon-ranking-2026.md §6):
  - до 30 тегов на карточку, каждый до 30 знаков;
  - только кириллица, латиница, цифры и «_»;
  - начинается с «#»;
  - не дублировать бренд, категорию и характеристики.

Последнее правилами не проверить машинно во всех случаях, но очевидный
случай ловим: тег не должен повторять название своей категории.

Запуск:
    python3 scripts/ozon_check_tags.py <хештеги.tsv> [книга.xlsx]
"""

import re
import sys

MAX_TAGS = 30
MAX_TAG_LEN = 30
RECOMMENDED_MAX = 10
TAG_RE = re.compile(r"^#[А-Яа-яЁёA-Za-z0-9_]+$")


def check(category, tags):
    problems, warnings = [], []

    if len(tags) > MAX_TAGS:
        problems.append(f"тегов {len(tags)} > {MAX_TAGS}")
    elif len(tags) > RECOMMENDED_MAX:
        warnings.append(f"тегов {len(tags)}, рекомендуют до {RECOMMENDED_MAX}")

    for t in tags:
        if not TAG_RE.match(t):
            problems.append(f"недопустимый тег: {t}")
        if len(t) > MAX_TAG_LEN:
            problems.append(f"тег длиннее {MAX_TAG_LEN}: {t}")

    # Дубль категории: тег из тех же слов, что и название категории.
    cat_words = set(re.findall(r"[А-Яа-яЁёA-Za-z]+", category.lower()))
    for t in tags:
        tag_words = set(re.findall(r"[А-Яа-яЁёA-Za-z]+", t.lower()))
        if tag_words and tag_words <= cat_words:
            problems.append(f"дублирует категорию: {t}")

    if len(set(tags)) != len(tags):
        problems.append("есть повторяющиеся теги")

    return problems, warnings


def load(path):
    out = []
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        if not line or line.startswith("#") or "|" not in line:
            continue
        category, tags = (p.strip() for p in line.split("|", 1))
        out.append((category, tags.split()))
    return out


def put_in_book(rows, book_path):
    import openpyxl
    from openpyxl.styles import PatternFill

    wb = openpyxl.load_workbook(book_path)
    ws = wb["Товары"]
    col_cat = col_tags = None
    for c in ws[1]:
        if c.value == "Категория *":
            col_cat = c.column
        elif c.value == "#Хештеги":
            col_tags = c.column
    if not col_cat or not col_tags:
        sys.exit("в книге не нашлись колонки «Категория» и «#Хештеги»")

    by_cat = {cat: " ".join(tags) for cat, tags in rows}
    green = PatternFill("solid", fgColor="E2EFDA")
    filled = 0
    for r in range(3, ws.max_row + 1):
        cat = ws.cell(row=r, column=col_cat).value
        if cat in by_cat:
            ws.cell(row=r, column=col_tags, value=by_cat[cat]).fill = green
            filled += 1
    wb.save(book_path)
    return filled


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    rows = load(sys.argv[1])

    bad = 0
    for category, tags in rows:
        problems, warnings = check(category, tags)
        mark = "ОШИБКА" if problems else ("  ~   " if warnings else "  ok  ")
        if problems:
            bad += 1
        print(f"{mark} {category:<30} тегов {len(tags)}")
        for p in problems:
            print(f"         └─ {p}")
        for w in warnings:
            print(f"         └─ {w}")

    print(f"\nкатегорий {len(rows)}, с ошибками {bad}")
    if bad:
        sys.exit(1)

    if len(sys.argv) > 2:
        n = put_in_book(rows, sys.argv[2])
        print(f"строк товаров с хештегами: {n}")
